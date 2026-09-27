#!/usr/bin/env python3
"""Create the formatted Word closure report from docs/closure-report.md."""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "closure-report.md"
OUTPUT = ROOT / "docs" / "Informe_de_cierre_EcoRed.docx"
GREEN = "176B50"
DARK = "17332B"
LIGHT_GREEN = "E8F3ED"


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_rich_text(paragraph, text: str) -> None:
    for segment in re.split(r"(\*\*.*?\*\*)", text):
        if segment.startswith("**") and segment.endswith("**"):
            paragraph.add_run(segment[2:-2]).bold = True
        elif segment:
            paragraph.add_run(segment)


def table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        values = [value.strip() for value in line.strip().strip("|").split("|")]
        if values and all(re.fullmatch(r":?-{3,}:?", value) for value in values):
            continue
        rows.append(values)
    return rows


def add_table(document, lines: list[str]) -> None:
    rows = table_rows(lines)
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_number, row_data in enumerate(rows):
        cells = table.add_row().cells
        for column_number, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = row_data[column_number] if column_number < len(row_data) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(3)
            add_rich_text(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if row_number == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if row_number == 0:
                shade_cell(cell, GREEN)
            elif row_number % 2 == 0:
                shade_cell(cell, LIGHT_GREEN)

    document.add_paragraph()


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def build_report() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(7)

    for style_name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(GREEN)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("ECORED  |  INFORME DE CIERRE").font.color.rgb = RGBColor.from_string(GREEN)
    header.runs[0].font.size = Pt(8)
    add_page_number(section.footer.paragraphs[0])

    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(110)
    title = cover.add_run("EcoRed")
    title.font.name = "Aptos Display"
    title.font.size = Pt(36)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(GREEN)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(12)
    subtitle_run = subtitle.add_run("INFORME DE CIERRE DEL PROYECTO")
    subtitle_run.font.name = "Aptos Display"
    subtitle_run.font.size = Pt(18)
    subtitle_run.font.bold = True
    subtitle_run.font.color.rgb = RGBColor.from_string(DARK)

    description = document.add_paragraph()
    description.alignment = WD_ALIGN_PARAGRAPH.CENTER
    description.paragraph_format.space_before = Pt(8)
    description.add_run("Donaciones de materiales y conexión con organizaciones receptoras").italic = True

    for detail in (
        "Luis Fernando Núñez Díaz",
        "Materia: Ingeniería de Software",
        "Docente: Luis Carlos Tarango Colmenero",
        "26 de septiembre de 2026",
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.add_run(detail)

    document.add_page_break()
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    skipped_main_title = False
    while index < len(lines):
        line = lines[index].strip()
        if not line or line == "---":
            index += 1
            continue

        if line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index].strip())
                index += 1
            add_table(document, block)
            continue

        if line.startswith("# "):
            if not skipped_main_title:
                skipped_main_title = True
            else:
                document.add_heading(line[2:].strip(), level=1)
            index += 1
            continue

        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
            index += 1
            continue

        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
            index += 1
            continue

        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_rich_text(paragraph, line[2:])
            index += 1
            continue

        if re.match(r"^\d+\.\s", line):
            paragraph = document.add_paragraph(style="List Number")
            add_rich_text(paragraph, re.sub(r"^\d+\.\s", "", line))
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_rich_text(paragraph, line)
        index += 1

    document.core_properties.title = "Informe de cierre del proyecto EcoRed"
    document.core_properties.subject = "Donaciones y reciclaje"
    document.core_properties.author = "EcoRed"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Informe guardado en {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    build_report()